import sqlite3
import json
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os

DB_DIR = os.path.dirname(os.path.abspath(__file__))
ORIGINAL_DB = os.path.join(DB_DIR, 'proveidos.db')

def get_db_path():
    if os.environ.get('VERCEL'):
        import shutil
        target = '/tmp/proveidos.db'
        if not os.path.exists(target) and os.path.exists(ORIGINAL_DB):
            shutil.copy2(ORIGINAL_DB, target)
        return target
    return ORIGINAL_DB

def get_db():
    conn = sqlite3.connect(get_db_path())
    conn.row_factory = sqlite3.Row
    return conn

def get_config(key, default=None):
    with get_db() as conn:
        row = conn.execute("SELECT valor FROM configuracion WHERE clave = ?", (key,)).fetchone()
        return row['valor'] if row else default

def set_config(key, value):
    with get_db() as conn:
        conn.execute("INSERT OR REPLACE INTO configuracion (clave, valor) VALUES (?, ?)", (key, str(value)))
        conn.commit()

def get_all_config():
    defaults = {
        'anio_activo': '2026',
        'institucion': 'UNIVERSIDAD NACIONAL HERMILIO VALDIZÁN',
        'dependencia': 'ESCUELA DE POSGRADO',
        'lugar_defecto': 'Cayhuayna',
        'sufijo_defecto': '-2026-UNHEVAL/EPG-D',
        'margen_top': '20',
        'margen_left': '22',
        'margen_right': '18',
        'margen_bottom': '18',
        'altura_firma': '140',
        'directora_nombre': 'Dra. Digna Amabilia Manrique de Lara Suarez',
        'directora_cargo': 'DIRECTORA'
    }
    with get_db() as conn:
        rows = conn.execute("SELECT clave, valor FROM configuracion").fetchall()
        for r in rows:
            defaults[r['clave']] = r['valor']
    return defaults

def save_config_dict(data):
    with get_db() as conn:
        for k, v in data.items():
            conn.execute("INSERT OR REPLACE INTO configuracion (clave, valor) VALUES (?, ?)", (k, str(v)))
        conn.commit()
    return get_all_config()

def get_next_proveido_numero(anio=2026):
    with get_db() as conn:
        row = conn.execute("SELECT MAX(numero) as max_num FROM proveidos WHERE anio = ?", (anio,)).fetchone()
        max_num = row['max_num'] if row and row['max_num'] is not None else 0
        return max_num + 1

def get_next_elevacion_numero(anio=2026):
    with get_db() as conn:
        row = conn.execute("SELECT MAX(numero) as max_num FROM elevaciones WHERE anio = ?", (anio,)).fetchone()
        max_num = row['max_num'] if row and row['max_num'] is not None else 0
        return max_num + 1

def get_proveidos(query="", fecha_desde="", fecha_hasta="", estado_impresion="todos", page=1, per_page=25):
    offset = (page - 1) * per_page
    where_clauses = ["(eliminado = 0 OR eliminado IS NULL)"]
    params = []

    if query:
        q_like = f"%{query.strip()}%"
        where_clauses.append("""
            (codigo_formateado LIKE ? OR
             destinatarios LIKE ? OR
             asunto LIKE ? OR
             referencia LIKE ? OR
             remitente_ref LIKE ? OR
             registro LIKE ?)
        """)
        params.extend([q_like, q_like, q_like, q_like, q_like, q_like])

    if fecha_desde:
        where_clauses.append("fecha_emision >= ?")
        params.append(fecha_desde)

    if fecha_hasta:
        where_clauses.append("fecha_emision <= ?")
        params.append(fecha_hasta)

    if estado_impresion == "pendientes":
        where_clauses.append("(impreso = 0 OR impreso IS NULL)")
    elif estado_impresion == "impresos":
        where_clauses.append("impreso = 1")

    where_sql = " WHERE " + " AND ".join(where_clauses)

    with get_db() as conn:
        count_sql = f"SELECT COUNT(*) as total FROM proveidos{where_sql}"
        total = conn.execute(count_sql, params).fetchone()['total']

        sql = f"""
            SELECT * FROM proveidos
            {where_sql}
            ORDER BY numero DESC
            LIMIT ? OFFSET ?
        """
        rows = conn.execute(sql, params + [per_page, offset]).fetchall()
        
        items = []
        for r in rows:
            d = dict(r)
            try:
                d['destinatarios'] = json.loads(d['destinatarios'])
            except:
                d['destinatarios'] = [{"nombre": d.get('destinatario_principal', ''), "cargo": d.get('cargo_principal', '')}]
            try:
                d['check_opciones'] = json.loads(d.get('check_opciones') or '[]')
            except:
                d['check_opciones'] = []
            d['veces_impreso'] = d.get('veces_impreso') or 0
            d['impreso'] = 1 if d.get('impreso') else 0
            items.append(d)

        # Conteo de pendientes para el badge
        total_pendientes = conn.execute("SELECT COUNT(*) as c FROM proveidos WHERE (impreso = 0 OR impreso IS NULL) AND (eliminado = 0 OR eliminado IS NULL)").fetchone()['c']

        return {
            "items": items,
            "total": total,
            "total_pendientes": total_pendientes,
            "page": page,
            "per_page": per_page,
            "total_pages": (total + per_page - 1) // per_page if total > 0 else 1
        }

def get_proveido_by_id(proveido_id):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM proveidos WHERE id = ?", (proveido_id,)).fetchone()
        if not row:
            return None
        d = dict(row)
        try:
            d['destinatarios'] = json.loads(d['destinatarios'])
        except:
            d['destinatarios'] = [{"nombre": d.get('destinatario_principal', ''), "cargo": d.get('cargo_principal', '')}]
        try:
            d['check_opciones'] = json.loads(d.get('check_opciones') or '[]')
        except:
            d['check_opciones'] = []
        d['veces_impreso'] = d.get('veces_impreso') or 0
        d['impreso'] = 1 if d.get('impreso') else 0
        return d

def marcar_proveido_impreso(proveido_id):
    now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    with get_db() as conn:
        conn.execute("""
            UPDATE proveidos SET
                impreso = 1,
                veces_impreso = COALESCE(veces_impreso, 0) + 1,
                fecha_impresion = ?
            WHERE id = ?
        """, (now_str, proveido_id))
        conn.commit()

def marcar_proveidos_impresos_lote(ids):
    if not ids:
        return
    now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    with get_db() as conn:
        for pid in ids:
            conn.execute("""
                UPDATE proveidos SET
                    impreso = 1,
                    veces_impreso = COALESCE(veces_impreso, 0) + 1,
                    fecha_impresion = ?
                WHERE id = ?
            """, (now_str, pid))
        conn.commit()

def toggle_impreso_proveido(proveido_id):
    with get_db() as conn:
        row = conn.execute("SELECT impreso, veces_impreso FROM proveidos WHERE id = ?", (proveido_id,)).fetchone()
        if not row:
            return False
        nuevo_impreso = 0 if row['impreso'] == 1 else 1
        nuevo_veces = row['veces_impreso']
        if nuevo_impreso == 1 and nuevo_veces == 0:
            nuevo_veces = 1
        elif nuevo_impreso == 0:
            nuevo_veces = 0
            
        now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S') if nuevo_impreso == 1 else None
        conn.execute("""
            UPDATE proveidos SET
                impreso = ?,
                veces_impreso = ?,
                fecha_impresion = ?
            WHERE id = ?
        """, (nuevo_impreso, nuevo_veces, now_str, proveido_id))
        conn.commit()
        return {"impreso": nuevo_impreso, "veces_impreso": nuevo_veces}

def get_proveidos_by_range(num_desde, num_hasta):
    # Sanitize inputs: jamás negativos
    num_desde = max(1, int(num_desde))
    num_hasta = max(num_desde, int(num_hasta))

    with get_db() as conn:
        rows = conn.execute("""
            SELECT * FROM proveidos
            WHERE numero >= ? AND numero <= ? AND (eliminado = 0 OR eliminado IS NULL)
            ORDER BY numero ASC
        """, (num_desde, num_hasta)).fetchall()
        
        items = []
        now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        for r in rows:
            d = dict(r)
            try:
                d['destinatarios'] = json.loads(d['destinatarios'])
            except:
                d['destinatarios'] = [{"nombre": d.get('destinatario_principal', ''), "cargo": d.get('cargo_principal', '')}]
            try:
                d['check_opciones'] = json.loads(d.get('check_opciones') or '[]')
            except:
                d['check_opciones'] = []
            d['veces_impreso'] = (d.get('veces_impreso') or 0) + 1
            d['impreso'] = 1
            # Auto-mark as printed
            conn.execute("""
                UPDATE proveidos SET
                    impreso = 1,
                    veces_impreso = COALESCE(veces_impreso, 0) + 1,
                    fecha_impresion = ?
                WHERE id = ?
            """, (now_str, d['id']))
            items.append(d)
        conn.commit()
        return items

def get_rango_info():
    with get_db() as conn:
        row = conn.execute("SELECT MIN(numero) as min_num, MAX(numero) as max_num, COUNT(*) as total FROM proveidos WHERE (eliminado = 0 OR eliminado IS NULL)").fetchone()
        pendientes = conn.execute("SELECT MIN(numero) as min_pend, MAX(numero) as max_pend, COUNT(*) as count_pend FROM proveidos WHERE (impreso = 0 OR impreso IS NULL) AND (eliminado = 0 OR eliminado IS NULL)").fetchone()
        numeros = [r['numero'] for r in conn.execute("SELECT numero FROM proveidos WHERE (eliminado = 0 OR eliminado IS NULL) ORDER BY numero DESC").fetchall()]
        return {
            "min_num": row['min_num'] or 1,
            "max_num": row['max_num'] or 1,
            "total": row['total'] or 0,
            "pendientes_min": pendientes['min_pend'],
            "pendientes_max": pendientes['max_pend'],
            "pendientes_count": pendientes['count_pend'] or 0,
            "numeros": numeros
        }

def verificar_reg_duplicado(reg_str):
    if not reg_str or not reg_str.strip():
        return {"existe": False}
    reg_clean = reg_str.strip()
    with get_db() as conn:
        row = conn.execute("""
            SELECT id, numero, codigo_formateado, destinatario_principal as destinatario, asunto, fecha_emision
            FROM proveidos
            WHERE registro = ? AND (eliminado = 0 OR eliminado IS NULL)
            ORDER BY numero DESC LIMIT 1
        """, (reg_clean,)).fetchone()
        if row:
            return {
                "existe": True,
                "tipo": "Proveído",
                "id": row['id'],
                "numero": row['codigo_formateado'],
                "destinatario": row['destinatario'],
                "asunto": row['asunto'],
                "fecha": row['fecha_emision']
            }
        row_e = conn.execute("""
            SELECT id, numero, codigo_formateado, destinatario, asunto, fecha_emision
            FROM elevaciones
            WHERE registro = ? AND (eliminado = 0 OR eliminado IS NULL)
            ORDER BY numero DESC LIMIT 1
        """, (reg_clean,)).fetchone()
        if row_e:
            return {
                "existe": True,
                "tipo": "Elevación",
                "id": row_e['id'],
                "numero": row_e['codigo_formateado'],
                "destinatario": row_e['destinatario'],
                "asunto": row_e['asunto'],
                "fecha": row_e['fecha_emision']
            }
        return {"existe": False}

def save_proveido(data):
    # Validate positive numbers
    numero = max(1, int(data.get('numero') or get_next_proveido_numero()))
    anio = max(2020, int(data.get('anio') or 2026))
    cod_formateado = f"{numero:04d}" if numero < 1000 else f"{numero}"
    sufijo = data.get('sufijo') or f"-{anio}-UNHEVAL/EPG-D"
    
    destinatarios = data.get('destinatarios', [])
    if isinstance(destinatarios, str):
        try:
            destinatarios = json.loads(destinatarios)
        except:
            destinatarios = [{"nombre": destinatarios, "cargo": data.get('cargo_principal', '')}]
            
    if not destinatarios:
        destinatarios = [{"nombre": data.get('destinatario_principal', ''), "cargo": data.get('cargo_principal', '')}]

    tipo = "MULTIPLE" if len(destinatarios) > 1 else "SIMPLE"
    dest_principal = destinatarios[0].get('nombre', '').strip() if destinatarios else ""
    cargo_principal = destinatarios[0].get('cargo', '').strip() if destinatarios else ""
    
    check_opciones = data.get('check_opciones', [])
    if isinstance(check_opciones, str):
        try:
            check_opciones = json.loads(check_opciones)
        except:
            check_opciones = []

    proveido_id = data.get('id')

    with get_db() as conn:
        for d in destinatarios:
            nom = (d.get('nombre') or '').strip()
            car = (d.get('cargo') or '').strip()
            if nom:
                conn.execute("""
                    INSERT INTO destinatarios_frecuentes (nombre, cargo, frecuencia)
                    VALUES (?, ?, 1)
                    ON CONFLICT(nombre) DO UPDATE SET
                    cargo = CASE WHEN ? != '' THEN ? ELSE cargo END,
                    frecuencia = frecuencia + 1
                """, (nom, car, car, car))

        if proveido_id:
            conn.execute("""
                UPDATE proveidos SET
                    numero = ?, codigo_formateado = ?, anio = ?, sufijo = ?,
                    tipo = ?, destinatarios = ?, destinatario_principal = ?, cargo_principal = ?,
                    asunto = ?, referencia = ?, remitente_ref = ?, registro = ?,
                    fecha_recibido = ?, fecha_emision = ?, lugar_emision = ?,
                    disposicion = ?, check_opciones = ?, dni = ?, titulo_art = ?, estado = ?,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
            """, (
                numero, cod_formateado, anio, sufijo,
                tipo, json.dumps(destinatarios, ensure_ascii=False), dest_principal, cargo_principal,
                data.get('asunto', '').strip(), data.get('referencia', '').strip(),
                data.get('remitente_ref', '').strip(), data.get('registro', '').strip(),
                data.get('fecha_recibido', ''), data.get('fecha_emision', ''),
                data.get('lugar_emision', 'Cayhuayna'), data.get('disposicion', '').strip(),
                json.dumps(check_opciones, ensure_ascii=False), data.get('dni', ''),
                data.get('titulo_art', ''), data.get('estado', 'EMITIDO'),
                proveido_id
            ))
            pid = proveido_id
        else:
            cur = conn.execute("""
                INSERT INTO proveidos (
                    numero, codigo_formateado, anio, sufijo,
                    tipo, destinatarios, destinatario_principal, cargo_principal,
                    asunto, referencia, remitente_ref, registro,
                    fecha_recibido, fecha_emision, lugar_emision,
                    disposicion, check_opciones, dni, titulo_art, estado,
                    impreso, veces_impreso
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, 0)
            """, (
                numero, cod_formateado, anio, sufijo,
                tipo, json.dumps(destinatarios, ensure_ascii=False), dest_principal, cargo_principal,
                data.get('asunto', '').strip(), data.get('referencia', '').strip(),
                data.get('remitente_ref', '').strip(), data.get('registro', '').strip(),
                data.get('fecha_recibido', ''), data.get('fecha_emision', ''),
                data.get('lugar_emision', 'Cayhuayna'), data.get('disposicion', '').strip(),
                json.dumps(check_opciones, ensure_ascii=False), data.get('dni', ''),
                data.get('titulo_art', ''), data.get('estado', 'EMITIDO')
            ))
            pid = cur.lastrowid
        conn.commit()
    return pid

def delete_proveido(proveido_id):
    # Enviar a papelera de reciclaje (soft delete)
    now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    with get_db() as conn:
        conn.execute("UPDATE proveidos SET eliminado = 1, fecha_eliminacion = ? WHERE id = ?", (now_str, proveido_id))
        conn.commit()
        return True

def restaurar_proveido(proveido_id):
    with get_db() as conn:
        conn.execute("UPDATE proveidos SET eliminado = 0, fecha_eliminacion = NULL WHERE id = ?", (proveido_id,))
        conn.commit()
        return True

def eliminar_permanente_proveido(proveido_id):
    with get_db() as conn:
        conn.execute("DELETE FROM proveidos WHERE id = ?", (proveido_id,))
        conn.commit()
        return True

def get_elevaciones(query="", fecha_desde="", fecha_hasta="", estado_impresion="todos", page=1, per_page=25):
    offset = (page - 1) * per_page
    where_clauses = ["(eliminado = 0 OR eliminado IS NULL)"]
    params = []

    if query:
        q_like = f"%{query.strip()}%"
        where_clauses.append("""
            (codigo_formateado LIKE ? OR
             destinatario LIKE ? OR
             asunto LIKE ? OR
             referencia LIKE ? OR
             remitente_ref LIKE ? OR
             registro LIKE ?)
        """)
        params.extend([q_like, q_like, q_like, q_like, q_like, q_like])

    if fecha_desde:
        where_clauses.append("fecha_emision >= ?")
        params.append(fecha_desde)

    if fecha_hasta:
        where_clauses.append("fecha_emision <= ?")
        params.append(fecha_hasta)

    if estado_impresion == "pendientes":
        where_clauses.append("(impreso = 0 OR impreso IS NULL)")
    elif estado_impresion == "impresos":
        where_clauses.append("impreso = 1")

    where_sql = " WHERE " + " AND ".join(where_clauses)

    with get_db() as conn:
        count_sql = f"SELECT COUNT(*) as total FROM elevaciones{where_sql}"
        total = conn.execute(count_sql, params).fetchone()['total']

        sql = f"""
            SELECT * FROM elevaciones
            {where_sql}
            ORDER BY numero DESC
            LIMIT ? OFFSET ?
        """
        rows = conn.execute(sql, params + [per_page, offset]).fetchall()
        
        items = []
        for r in rows:
            d = dict(r)
            try:
                d['check_opciones'] = json.loads(d.get('check_opciones') or '[]')
            except:
                d['check_opciones'] = []
            d['veces_impreso'] = d.get('veces_impreso') or 0
            d['impreso'] = 1 if d.get('impreso') else 0
            items.append(d)

        # Conteo de pendientes para el badge
        total_pendientes = conn.execute("SELECT COUNT(*) as c FROM elevaciones WHERE (impreso = 0 OR impreso IS NULL) AND (eliminado = 0 OR eliminado IS NULL)").fetchone()['c']

        return {
            "items": items,
            "total": total,
            "total_pendientes": total_pendientes,
            "page": page,
            "per_page": per_page,
            "total_pages": (total + per_page - 1) // per_page if total > 0 else 1
        }

def get_elevacion_by_id(elevacion_id):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM elevaciones WHERE id = ?", (elevacion_id,)).fetchone()
        if not row:
            return None
        d = dict(row)
        try:
            d['check_opciones'] = json.loads(d.get('check_opciones') or '[]')
        except:
            d['check_opciones'] = []
        d['veces_impreso'] = d.get('veces_impreso') or 0
        d['impreso'] = 1 if d.get('impreso') else 0
        return d

def marcar_elevacion_impreso(elevacion_id):
    now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    with get_db() as conn:
        conn.execute("""
            UPDATE elevaciones SET
                impreso = 1,
                veces_impreso = COALESCE(veces_impreso, 0) + 1,
                fecha_impresion = ?
            WHERE id = ?
        """, (now_str, elevacion_id))
        conn.commit()

def marcar_elevaciones_impresas_lote(ids):
    if not ids:
        return
    now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    with get_db() as conn:
        for eid in ids:
            conn.execute("""
                UPDATE elevaciones SET
                    impreso = 1,
                    veces_impreso = COALESCE(veces_impreso, 0) + 1,
                    fecha_impresion = ?
                WHERE id = ?
            """, (now_str, eid))
        conn.commit()

def toggle_impreso_elevacion(elevacion_id):
    with get_db() as conn:
        row = conn.execute("SELECT impreso, veces_impreso FROM elevaciones WHERE id = ?", (elevacion_id,)).fetchone()
        if not row:
            return False
        nuevo_impreso = 0 if row['impreso'] == 1 else 1
        nuevo_veces = row['veces_impreso']
        if nuevo_impreso == 1 and nuevo_veces == 0:
            nuevo_veces = 1
        elif nuevo_impreso == 0:
            nuevo_veces = 0
            
        now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S') if nuevo_impreso == 1 else None
        conn.execute("""
            UPDATE elevaciones SET
                impreso = ?,
                veces_impreso = ?,
                fecha_impresion = ?
            WHERE id = ?
        """, (nuevo_impreso, nuevo_veces, now_str, elevacion_id))
        conn.commit()
        return {"impreso": nuevo_impreso, "veces_impreso": nuevo_veces}

def get_elevaciones_by_range(num_desde, num_hasta):
    # Sanitize inputs: jamás negativos
    num_desde = max(1, int(num_desde))
    num_hasta = max(num_desde, int(num_hasta))

    with get_db() as conn:
        rows = conn.execute("""
            SELECT * FROM elevaciones
            WHERE numero >= ? AND numero <= ? AND (eliminado = 0 OR eliminado IS NULL)
            ORDER BY numero ASC
        """, (num_desde, num_hasta)).fetchall()
        
        items = []
        now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        for r in rows:
            d = dict(r)
            try:
                d['check_opciones'] = json.loads(d.get('check_opciones') or '[]')
            except:
                d['check_opciones'] = []
            d['veces_impreso'] = (d.get('veces_impreso') or 0) + 1
            d['impreso'] = 1
            # Auto-mark as printed
            conn.execute("""
                UPDATE elevaciones SET
                    impreso = 1,
                    veces_impreso = COALESCE(veces_impreso, 0) + 1,
                    fecha_impresion = ?
                WHERE id = ?
            """, (now_str, d['id']))
            items.append(d)
        conn.commit()
        return items

def get_rango_info_elevaciones():
    with get_db() as conn:
        row = conn.execute("SELECT MIN(numero) as min_num, MAX(numero) as max_num, COUNT(*) as total FROM elevaciones WHERE (eliminado = 0 OR eliminado IS NULL)").fetchone()
        pendientes = conn.execute("SELECT MIN(numero) as min_pend, MAX(numero) as max_pend, COUNT(*) as count_pend FROM elevaciones WHERE (impreso = 0 OR impreso IS NULL) AND (eliminado = 0 OR eliminado IS NULL)").fetchone()
        numeros = [r['numero'] for r in conn.execute("SELECT numero FROM elevaciones WHERE (eliminado = 0 OR eliminado IS NULL) ORDER BY numero DESC").fetchall()]
        return {
            "min_num": row['min_num'] or 1,
            "max_num": row['max_num'] or 1,
            "total": row['total'] or 0,
            "pendientes_min": pendientes['min_pend'],
            "pendientes_max": pendientes['max_pend'],
            "pendientes_count": pendientes['count_pend'] or 0,
            "numeros": numeros
        }

def save_elevacion(data):
    numero = max(1, int(data.get('numero') or get_next_elevacion_numero()))
    anio = max(2020, int(data.get('anio') or 2026))
    cod_formateado = f"{numero:04d}"
    sufijo = data.get('sufijo') or f"-{anio}-UNHEVAL/EPG-D"
    check_opciones = data.get('check_opciones', [])
    if isinstance(check_opciones, str):
        try:
            check_opciones = json.loads(check_opciones)
        except:
            check_opciones = []

    elevacion_id = data.get('id')

    with get_db() as conn:
        if elevacion_id:
            conn.execute("""
                UPDATE elevaciones SET
                    numero = ?, codigo_formateado = ?, anio = ?, sufijo = ?,
                    destinatario = ?, asunto = ?, referencia = ?, remitente_ref = ?,
                    registro = ?, fecha_recibido = ?, fecha_emision = ?,
                    lugar_emision = ?, disposicion = ?, check_opciones = ?, doc_ref = ?,
                    estado = ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
            """, (
                numero, cod_formateado, anio, sufijo,
                data.get('destinatario', '').strip(), data.get('asunto', '').strip(),
                data.get('referencia', '').strip(), data.get('remitente_ref', '').strip(),
                data.get('registro', '').strip(), data.get('fecha_recibido', ''),
                data.get('fecha_emision', ''), data.get('lugar_emision', 'Cayhuayna'),
                data.get('disposicion', '').strip(), json.dumps(check_opciones, ensure_ascii=False),
                data.get('doc_ref', '').strip(), data.get('estado', 'EMITIDO'),
                elevacion_id
            ))
            eid = elevacion_id
        else:
            cur = conn.execute("""
                INSERT INTO elevaciones (
                    numero, codigo_formateado, anio, sufijo,
                    destinatario, asunto, referencia, remitente_ref,
                    registro, fecha_recibido, fecha_emision,
                    lugar_emision, disposicion, check_opciones, doc_ref, estado,
                    impreso, veces_impreso
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, 0)
            """, (
                numero, cod_formateado, anio, sufijo,
                data.get('destinatario', '').strip(), data.get('asunto', '').strip(),
                data.get('referencia', '').strip(), data.get('remitente_ref', '').strip(),
                data.get('registro', '').strip(), data.get('fecha_recibido', ''),
                data.get('fecha_emision', ''), data.get('lugar_emision', 'Cayhuayna'),
                data.get('disposicion', '').strip(), json.dumps(check_opciones, ensure_ascii=False),
                data.get('doc_ref', '').strip(), data.get('estado', 'EMITIDO')
            ))
            eid = cur.lastrowid
        conn.commit()
    return eid

def delete_elevacion(elevacion_id):
    # Enviar a papelera de reciclaje (soft delete)
    now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    with get_db() as conn:
        conn.execute("UPDATE elevaciones SET eliminado = 1, fecha_eliminacion = ? WHERE id = ?", (now_str, elevacion_id))
        conn.commit()
        return True

def restaurar_elevacion(elevacion_id):
    with get_db() as conn:
        conn.execute("UPDATE elevaciones SET eliminado = 0, fecha_eliminacion = NULL WHERE id = ?", (elevacion_id,))
        conn.commit()
        return True

def eliminar_permanente_elevacion(elevacion_id):
    with get_db() as conn:
        conn.execute("DELETE FROM elevaciones WHERE id = ?", (elevacion_id,))
        conn.commit()
        return True

def purgar_papelera():
    # Purga automática de documentos con más de 30 días (1 mes) en papelera
    limite = (datetime.datetime.now() - datetime.timedelta(days=30)).strftime('%Y-%m-%d %H:%M:%S')
    with get_db() as conn:
        conn.execute("DELETE FROM proveidos WHERE eliminado = 1 AND fecha_eliminacion < ?", (limite,))
        conn.execute("DELETE FROM elevaciones WHERE eliminado = 1 AND fecha_eliminacion < ?", (limite,))
        conn.commit()

def vaciar_papelera():
    with get_db() as conn:
        conn.execute("DELETE FROM proveidos WHERE eliminado = 1")
        conn.execute("DELETE FROM elevaciones WHERE eliminado = 1")
        conn.commit()
        return True

def get_papelera():
    purgar_papelera()
    items = []
    now = datetime.datetime.now()
    with get_db() as conn:
        rows_p = conn.execute("""
            SELECT id, numero, codigo_formateado, anio, sufijo, asunto, destinatarios,
                   destinatario_principal, registro, fecha_emision, fecha_eliminacion
            FROM proveidos WHERE eliminado = 1
            ORDER BY fecha_eliminacion DESC
        """).fetchall()
        for r in rows_p:
            d = dict(r)
            d['tipo_doc'] = 'PROVEÍDO'
            dias_restantes = 30
            if d.get('fecha_eliminacion'):
                try:
                    fe = datetime.datetime.strptime(d['fecha_eliminacion'][:19], '%Y-%m-%d %H:%M:%S')
                    dias_pasados = (now - fe).days
                    dias_restantes = max(0, 30 - dias_pasados)
                except:
                    dias_restantes = 30
            d['dias_restantes'] = dias_restantes
            try:
                d['destinatarios'] = json.loads(d['destinatarios'])
            except:
                d['destinatarios'] = [{"nombre": d.get('destinatario_principal', '')}]
            items.append(d)

        rows_e = conn.execute("""
            SELECT id, numero, codigo_formateado, anio, sufijo, asunto, destinatario,
                   registro, fecha_emision, fecha_eliminacion
            FROM elevaciones WHERE eliminado = 1
            ORDER BY fecha_eliminacion DESC
        """).fetchall()
        for r in rows_e:
            d = dict(r)
            d['tipo_doc'] = 'ELEVACIÓN'
            dias_restantes = 30
            if d.get('fecha_eliminacion'):
                try:
                    fe = datetime.datetime.strptime(d['fecha_eliminacion'][:19], '%Y-%m-%d %H:%M:%S')
                    dias_pasados = (now - fe).days
                    dias_restantes = max(0, 30 - dias_pasados)
                except:
                    dias_restantes = 30
            d['dias_restantes'] = dias_restantes
            items.append(d)

    items.sort(key=lambda x: x.get('fecha_eliminacion') or '', reverse=True)
    return {
        "total": len(items),
        "items": items
    }

def get_destinatarios_sugerencias(term=""):
    with get_db() as conn:
        term_like = f"%{term.strip()}%" if term else "%"
        rows = conn.execute("""
            SELECT nombre, cargo, frecuencia
            FROM destinatarios_frecuentes
            WHERE nombre LIKE ? OR cargo LIKE ?
            ORDER BY frecuencia DESC, nombre ASC
            LIMIT 25
        """, (term_like, term_like)).fetchall()
        return [dict(r) for r in rows]

def get_dashboard_analytics():
    with get_db() as conn:
        total_prov = conn.execute("SELECT COUNT(*) as c FROM proveidos WHERE (eliminado = 0 OR eliminado IS NULL)").fetchone()['c']
        total_elev = conn.execute("SELECT COUNT(*) as c FROM elevaciones WHERE (eliminado = 0 OR eliminado IS NULL)").fetchone()['c']
        total_exp_p = conn.execute("SELECT DISTINCT registro FROM proveidos WHERE length(registro) > 0 AND (eliminado = 0 OR eliminado IS NULL)").fetchall()
        total_exp_e = conn.execute("SELECT DISTINCT registro FROM elevaciones WHERE length(registro) > 0 AND (eliminado = 0 OR eliminado IS NULL)").fetchall()
        all_regs = set([r['registro'] for r in total_exp_p] + [r['registro'] for r in total_exp_e])
        total_expedientes = len(all_regs)

        total_impresos_p = conn.execute("SELECT COUNT(*) as c FROM proveidos WHERE impreso = 1 AND (eliminado = 0 OR eliminado IS NULL)").fetchone()['c']
        total_impresos_e = conn.execute("SELECT COUNT(*) as c FROM elevaciones WHERE impreso = 1 AND (eliminado = 0 OR eliminado IS NULL)").fetchone()['c']
        total_impresos = total_impresos_p + total_impresos_e

        total_pendientes_p = conn.execute("SELECT COUNT(*) as c FROM proveidos WHERE (impreso = 0 OR impreso IS NULL) AND (eliminado = 0 OR eliminado IS NULL)").fetchone()['c']
        total_pendientes_e = conn.execute("SELECT COUNT(*) as c FROM elevaciones WHERE (impreso = 0 OR impreso IS NULL) AND (eliminado = 0 OR eliminado IS NULL)").fetchone()['c']
        total_pendientes = total_pendientes_p + total_pendientes_e

        meses_data = conn.execute("""
            SELECT substr(fecha_emision, 1, 7) as mes, count(*) as cantidad
            FROM proveidos
            WHERE length(fecha_emision) >= 7
            GROUP BY mes
            ORDER BY mes
        """).fetchall()

        meses_nombres = {
            '2026-01': 'Ene', '2026-02': 'Feb', '2026-03': 'Mar', '2026-04': 'Abr',
            '2026-05': 'May', '2026-06': 'Jun', '2026-07': 'Jul', '2026-08': 'Ago',
            '2026-09': 'Set', '2026-10': 'Oct', '2026-11': 'Nov', '2026-12': 'Dic'
        }
        
        labels_meses = []
        valores_meses = []
        for m in meses_data:
            if m['mes'] in meses_nombres:
                labels_meses.append(meses_nombres[m['mes']])
                valores_meses.append(m['cantidad'])

        top_dest = conn.execute("""
            SELECT destinatario_principal, count(*) as cantidad
            FROM proveidos
            WHERE length(destinatario_principal) > 0
            GROUP BY destinatario_principal
            ORDER BY cantidad DESC
            LIMIT 7
        """).fetchall()

        labels_dest = [r['destinatario_principal'][:24] + '...' if len(r['destinatario_principal']) > 24 else r['destinatario_principal'] for r in top_dest]
        valores_dest = [r['cantidad'] for r in top_dest]

        pagos_count = conn.execute("SELECT COUNT(*) as c FROM proveidos WHERE asunto LIKE '%PAGO%' OR asunto LIKE '%EVALUACION%' OR asunto LIKE '%ECONOMICA%'").fetchone()['c']
        tesis_count = conn.execute("SELECT COUNT(*) as c FROM proveidos WHERE asunto LIKE '%TESIS%' OR asunto LIKE '%JURADO%' OR asunto LIKE '%GRADO%'").fetchone()['c']
        clases_count = conn.execute("SELECT COUNT(*) as c FROM proveidos WHERE asunto LIKE '%MODALIDAD%' OR asunto LIKE '%CLASES%' OR asunto LIKE '%REPROGRAMACION%' OR asunto LIKE '%HORARIO%'").fetchone()['c']
        otros_count = max(0, total_prov - (pagos_count + tesis_count + clases_count))

        promedio_mes = round(total_prov / max(len(valores_meses), 1), 1)

        return {
            "total_proveidos": total_prov,
            "total_elevaciones": total_elev,
            "total_expedientes": total_expedientes,
            "total_impresos": total_impresos,
            "total_pendientes": total_pendientes,
            "promedio_mensual": promedio_mes,
            "labels_meses": labels_meses,
            "valores_meses": valores_meses,
            "labels_destinatarios": labels_dest,
            "valores_destinatarios": valores_dest,
            "categorias_tramite": {
                "Gestion_Economica_Pagos": pagos_count,
                "Tesis_y_Grados": tesis_count,
                "Docencia_y_Clases": clases_count,
                "Otros_Tramites": otros_count
            }
        }

def get_estadisticas():
    with get_db() as conn:
        total_prov = conn.execute("SELECT COUNT(*) as c FROM proveidos WHERE (eliminado = 0 OR eliminado IS NULL)").fetchone()['c']
        total_elev = conn.execute("SELECT COUNT(*) as c FROM elevaciones WHERE (eliminado = 0 OR eliminado IS NULL)").fetchone()['c']
        ultimo_prov = conn.execute("SELECT MAX(numero) as m FROM proveidos WHERE (eliminado = 0 OR eliminado IS NULL)").fetchone()['m'] or 0
        ultimo_elev = conn.execute("SELECT MAX(numero) as m FROM elevaciones WHERE (eliminado = 0 OR eliminado IS NULL)").fetchone()['m'] or 0
        
        ultimos_prov = conn.execute("""
            SELECT id, numero, codigo_formateado, destinatario_principal, asunto, fecha_emision
            FROM proveidos 
            WHERE (eliminado = 0 OR eliminado IS NULL)
            ORDER BY numero DESC LIMIT 5
        """).fetchall()
        
        return {
            "total_proveidos": total_prov,
            "total_elevaciones": total_elev,
            "ultimo_proveido_numero": ultimo_prov,
            "ultimo_elevacion_numero": ultimo_elev,
            "siguiente_proveido": ultimo_prov + 1,
            "siguiente_elevacion": ultimo_elev + 1,
            "ultimos_proveidos": [dict(r) for r in ultimos_prov]
        }

def export_to_excel(output_filename="EXPORT_PROVEIDOS_Y_ELEVACIONES_2026.xlsx"):
    wb = openpyxl.Workbook()
    ws_p = wb.active
    ws_p.title = "PROVEIDOS 2026"
    
    headers_p = ["N° PROV", "TIPO", "DESTINATARIOS / NOMBRES", "CARGO(S)", "ASUNTO", "REF.", "TA / REMITENTE", "REG", "FECHA RECIBIDO", "FECHA EMISION", "DISPOSICION", "IMPRESO", "VECES IMPRESO", "FECHA IMPRESION"]
    ws_p.append(headers_p)
    
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    
    for col_idx in range(1, len(headers_p) + 1):
        cell = ws_p.cell(1, col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    with get_db() as conn:
        rows_p = conn.execute("SELECT * FROM proveidos WHERE (eliminado = 0 OR eliminado IS NULL) ORDER BY numero ASC").fetchall()
        for r in rows_p:
            d = dict(r)
            try:
                dest_list = json.loads(d['destinatarios'])
                nombres = "\n".join([x.get('nombre', '') for x in dest_list if x.get('nombre')])
                cargos = "\n".join([x.get('cargo', '') for x in dest_list if x.get('cargo')])
            except:
                nombres = d.get('destinatario_principal', '')
                cargos = d.get('cargo_principal', '')
                
            ws_p.append([
                d.get('codigo_formateado'),
                d.get('tipo'),
                nombres,
                cargos,
                d.get('asunto'),
                d.get('referencia'),
                d.get('remitente_ref'),
                d.get('registro'),
                d.get('fecha_recibido'),
                d.get('fecha_emision'),
                d.get('disposicion'),
                "SI" if d.get('impreso') else "NO",
                d.get('veces_impreso') or 0,
                d.get('fecha_impresion') or ""
            ])

    ws_e = wb.create_sheet(title="ELEVACIONES 2026")
    headers_e = ["N° ELEVACION", "DESTINATARIO", "ASUNTO", "REF.", "REMITENTE / COMISION", "REG", "FECHA RECIBIDO", "FECHA EMISION", "DISPOSICION", "DOC REF"]
    ws_e.append(headers_e)
    
    for col_idx in range(1, len(headers_e) + 1):
        cell = ws_e.cell(1, col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    with get_db() as conn:
        rows_e = conn.execute("SELECT * FROM elevaciones WHERE (eliminado = 0 OR eliminado IS NULL) ORDER BY numero ASC").fetchall()
        for r in rows_e:
            d = dict(r)
            ws_e.append([
                d.get('codigo_formateado'),
                d.get('destinatario'),
                d.get('asunto'),
                d.get('referencia'),
                d.get('remitente_ref'),
                d.get('registro'),
                d.get('fecha_recibido'),
                d.get('fecha_emision'),
                d.get('disposicion'),
                d.get('doc_ref')
            ])
            
    wb.save(output_filename)
    return output_filename

# ========================================================
# GESTIÓN DE PLANTILLAS PREDEFINIDAS Y PERSONALIZADAS
# ========================================================

def init_plantillas_table():
    with get_db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS plantillas (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                tipo_doc TEXT NOT NULL,
                nombre TEXT NOT NULL,
                icono TEXT DEFAULT '📋',
                asunto TEXT,
                disposicion TEXT,
                destinatarios TEXT,
                destinatario TEXT,
                check_opciones TEXT,
                es_predeterminada INTEGER DEFAULT 0,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        cnt = conn.execute("SELECT COUNT(*) as c FROM plantillas").fetchone()['c']
        if cnt == 0:
            seed_default_plantillas(conn)
        conn.commit()

def seed_default_plantillas(conn, force_tipo=None):
    prov_seeds = [
        {
            "tipo_doc": "proveido",
            "nombre": "Trámite de Pago / Gestión Económica",
            "icono": "💳",
            "asunto": "SOLICITO SE SIRVA REALIZAR LA EVALUACIÓN CORRESPONDIENTE Y, DE ESTAR CONFORME, PROCEDER SEGÚN EL TRÁMITE ESTABLECIDO.",
            "disposicion": "DISPONGASE LA EVALUACION INMEDIATA DE LOS DOCUMENTOS MENCIONADAS EN EL DOCUMENTO DE REFERENCIA A FIN DE CONTINUAR CON EL TRAMITE PARA EL PAGO CORRESPONDIENTE",
            "destinatarios": [{"nombre": "FERNANDO JOAQUÍN ALMANZA GARAY", "cargo": "RESPONSABLE DEL AREA DE GESTION ECONOMICA DE LA EPG"}],
            "check_opciones": ["evaluar", "informar_presupuesto"]
        },
        {
            "tipo_doc": "proveido",
            "nombre": "Borrador de Tesis",
            "icono": "🎓",
            "asunto": "LEVANTAMIENTO DE OBSERVACIONES DEL BORRADOR DE TESIS DEL GRADO DE MAESTRO.",
            "disposicion": "EMITIR EL INFORME DENTRO DE LOS DIEZ (10) DIAS CALENDARIOS DE HABER RECEPCIONADO EL PRESENTE PROVEÍDO",
            "destinatarios": [],
            "check_opciones": ["su_conocimiento", "emitir_dictamen"]
        },
        {
            "tipo_doc": "proveido",
            "nombre": "Sustentación de Tesis",
            "icono": "📜",
            "asunto": "LEVANTAMIENTO DE OBSERVACIONES DE ACTA DE SUSTENTACION DE TESIS DEL GRADO DE MAESTRO.",
            "disposicion": "EMITIR EL INFORME DENTRO DE LOS DIEZ (10) DIAS CALENDARIOS DE HABER RECEPCIONADO EL PRESENTE PROVEÍDO",
            "destinatarios": [],
            "check_opciones": ["su_conocimiento", "emitir_dictamen"]
        },
        {
            "tipo_doc": "proveido",
            "nombre": "Convalidación de Cursos",
            "icono": "🔄",
            "asunto": "SOLICITA CONVALIDACION DE CURSOS",
            "disposicion": "PARA SU CONOCIMIENTO Y LAS ACCIONES QUE CORRESPONDAN",
            "destinatarios": [],
            "check_opciones": ["su_conocimiento", "evaluar"]
        },
        {
            "tipo_doc": "proveido",
            "nombre": "Solicitud de Informe",
            "icono": "📋",
            "asunto": "SOLICITO INFORME CORRESPONDIENTE A SU DEPENDENCIA",
            "disposicion": "EMITIR EL INFORME DENTRO DEL PLAZO ESTABLECIDO",
            "destinatarios": [],
            "check_opciones": ["informar"]
        },
        {
            "tipo_doc": "proveido",
            "nombre": "Actividades de Locador",
            "icono": "🧑‍💼",
            "asunto": "INFORME DE ACTIVIDADES REALIZADAS DEL LOCADOR DE SERVICIOS",
            "disposicion": "DISPONGASE LA EVALUACION INMEDIATA PARA EL TRÁMITE DE PAGO CORRESPONDIENTE",
            "destinatarios": [{"nombre": "FERNANDO JOAQUÍN ALMANZA GARAY", "cargo": "RESPONSABLE DEL AREA DE GESTION ECONOMICA DE LA EPG"}],
            "check_opciones": ["evaluar", "informar_presupuesto"]
        },
        {
            "tipo_doc": "proveido",
            "nombre": "Cambio Modalidad Clases",
            "icono": "💻",
            "asunto": "SE AUTORIZA EL CAMBIO DE MODALIDAD DE CLASES PRESENCIAL A VIRTUAL",
            "disposicion": "CAMBIO DE MODALIDAD DE CLASE PRESENCIAL A VIRTUAL EN EL HORARIO HABITUAL, A PETICIÓN DEL DOCENTE",
            "destinatarios": [],
            "check_opciones": ["su_conocimiento", "su_atencion"]
        },
        {
            "tipo_doc": "proveido",
            "nombre": "Informe de Notas",
            "icono": "📊",
            "asunto": "SOLICITA INFORME DE NOTAS",
            "disposicion": "EMITIR EL INFORME CORRESPONDIENTE",
            "destinatarios": [],
            "check_opciones": ["informar"]
        }
    ]

    elev_seeds = [
        {
            "tipo_doc": "elevacion",
            "nombre": "Traslado y Convalidación",
            "icono": "🏛️",
            "asunto": "SU REVISIÓN EN CONSEJO DIRECTIVO: SE SOLICITA TRASLADO EXTERNO Y CONVALIDACIÓN DE CURSOS",
            "disposicion": "SE ELEVA AL CONSEJO DIRECTIVO DE LA ESCUELA DE POSGRADO PARA SU REVISIÓN, EVALUACIÓN Y EMISIÓN DE LA RESOLUCIÓN CORRESPONDIENTE.",
            "destinatario": "CONSEJO DIRECTIVO DE LA ESCUELA DE POSGRADO",
            "check_opciones": ["consejo_universitario", "emitir_resolucion"]
        },
        {
            "tipo_doc": "elevacion",
            "nombre": "Elevación Proyecto Tesis",
            "icono": "🎓",
            "asunto": "ELEVACIÓN DE PROYECTO DE TESIS Y DICTAMEN DE ASESOR PARA SU APROBACIÓN POR CONSEJO DIRECTIVO",
            "disposicion": "SE ELEVA AL CONSEJO DIRECTIVO DE LA ESCUELA DE POSGRADO PARA SU CONOCIMIENTO, EVALUACIÓN Y EMISIÓN DE LA RESOLUCIÓN CORRESPONDIENTE.",
            "destinatario": "CONSEJO DIRECTIVO DE LA ESCUELA DE POSGRADO",
            "check_opciones": ["consejo_universitario", "emitir_resolucion", "su_atencion"]
        },
        {
            "tipo_doc": "elevacion",
            "nombre": "Dictamen de Grado",
            "icono": "📜",
            "asunto": "ELEVACIÓN DE EXPEDIENTE DE DICTAMEN FAVORABLE DE GRADO ACADÉMICO PARA OTORGAMIENTO DE DIPLOMA",
            "disposicion": "SE ELEVA CON DICTAMEN DE CONFORMIDAD A FIN DE QUE SE DISPONGA LA EMISIÓN DE LA RESOLUCIÓN DE CONSEJO DIRECTIVO Y EL OTORGAMIENTO DEL GRADO RESPECTIVO.",
            "destinatario": "CONSEJO DIRECTIVO DE LA ESCUELA DE POSGRADO",
            "check_opciones": ["consejo_universitario", "emitir_resolucion", "fines_pertinentes"]
        },
        {
            "tipo_doc": "elevacion",
            "nombre": "Plan de Trabajo / Admisión",
            "icono": "📑",
            "asunto": "SOLICITO MODIFICATORIA / APROBACIÓN DEL PLAN DE TRABAJO DEL PROCESO DE ADMISIÓN",
            "disposicion": "SE ELEVA PARA SU CONOCIMIENTO, EVALUACIÓN DE DISPONIBILIDAD PRESUPUESTAL Y APROBACIÓN POR EL PLENO DEL CONSEJO.",
            "destinatario": "CONSEJO DIRECTIVO DE LA ESCUELA DE POSGRADO",
            "check_opciones": ["informar_presupuesto", "consejo_universitario", "evaluar"]
        },
        {
            "tipo_doc": "elevacion",
            "nombre": "Ratificación de Resolución",
            "icono": "⚖️",
            "asunto": "SOLICITO LA RATIFICACIÓN DE RESOLUCIÓN DE FACULTAD",
            "disposicion": "SE ELEVA PARA SU CONOCIMIENTO Y EMISIÓN DEL ACTO RESOLUTIVO CORRESPONDIENTE.",
            "destinatario": "CONSEJO DIRECTIVO DE LA ESCUELA DE POSGRADO",
            "check_opciones": ["consejo_universitario", "emitir_resolucion"]
        }
    ]

    if force_tipo is None or force_tipo == 'proveido':
        for p in prov_seeds:
            conn.execute("""
                INSERT INTO plantillas (tipo_doc, nombre, icono, asunto, disposicion, destinatarios, check_opciones, es_predeterminada)
                VALUES (?, ?, ?, ?, ?, ?, ?, 1)
            """, (p['tipo_doc'], p['nombre'], p['icono'], p['asunto'], p['disposicion'], json.dumps(p['destinatarios'], ensure_ascii=False), json.dumps(p['check_opciones'], ensure_ascii=False)))

    if force_tipo is None or force_tipo == 'elevacion':
        for e in elev_seeds:
            conn.execute("""
                INSERT INTO plantillas (tipo_doc, nombre, icono, asunto, disposicion, destinatario, check_opciones, es_predeterminada)
                VALUES (?, ?, ?, ?, ?, ?, ?, 1)
            """, (e['tipo_doc'], e['nombre'], e['icono'], e['asunto'], e['disposicion'], e['destinatario'], json.dumps(e['check_opciones'], ensure_ascii=False)))

def get_plantillas(tipo_doc='proveido'):
    init_plantillas_table()
    with get_db() as conn:
        rows = conn.execute("SELECT * FROM plantillas WHERE tipo_doc = ? ORDER BY es_predeterminada DESC, id ASC", (tipo_doc,)).fetchall()
        items = []
        for r in rows:
            d = dict(r)
            try:
                d['destinatarios'] = json.loads(d.get('destinatarios') or '[]')
            except:
                d['destinatarios'] = []
            try:
                d['check_opciones'] = json.loads(d.get('check_opciones') or '[]')
            except:
                d['check_opciones'] = []
            items.append(d)
        return items

def save_plantilla(data):
    init_plantillas_table()
    tipo_doc = data.get('tipo_doc', 'proveido')
    nombre = data.get('nombre', '').strip()
    if not nombre:
        raise ValueError("El nombre de la plantilla es obligatorio")
    icono = data.get('icono', '📋').strip() or '📋'
    asunto = data.get('asunto', '').strip()
    disposicion = data.get('disposicion', '').strip()
    destinatarios = data.get('destinatarios', [])
    destinatario = data.get('destinatario', '').strip()
    check_opciones = data.get('check_opciones', [])

    with get_db() as conn:
        cur = conn.execute("""
            INSERT INTO plantillas (tipo_doc, nombre, icono, asunto, disposicion, destinatarios, destinatario, check_opciones, es_predeterminada)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, 0)
        """, (
            tipo_doc, nombre, icono, asunto, disposicion,
            json.dumps(destinatarios, ensure_ascii=False) if isinstance(destinatarios, list) else destinatarios,
            destinatario,
            json.dumps(check_opciones, ensure_ascii=False) if isinstance(check_opciones, list) else check_opciones
        ))
        conn.commit()
        return cur.lastrowid

def delete_plantilla(plantilla_id):
    with get_db() as conn:
        conn.execute("DELETE FROM plantillas WHERE id = ?", (plantilla_id,))
        conn.commit()
        return True

def reset_default_plantillas(tipo_doc='proveido'):
    init_plantillas_table()
    with get_db() as conn:
        conn.execute("DELETE FROM plantillas WHERE tipo_doc = ? AND es_predeterminada = 1", (tipo_doc,))
        seed_default_plantillas(conn, force_tipo=tipo_doc)
        conn.commit()
        return True


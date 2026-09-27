import openpyxl
import sqlite3
import json
import datetime
import os
import re

EXCEL_PATH = r'C:\Users\joell\Downloads\ELEVACION y PROVEIDO 2026 macros (1).xlsm'
DB_PATH = 'proveidos.db'

def format_date(val):
    if val is None:
        return ""
    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.strftime('%Y-%m-%d')
    s = str(val).strip()
    return s[:10] if len(s) >= 10 else s

def clean_str(val):
    if val is None:
        return ""
    return str(val).strip()

def parse_destinatarios(nombre_val, cargo_val):
    nombres_raw = clean_str(nombre_val)
    cargos_raw = clean_str(cargo_val)
    
    nombres_list = [n.strip() for n in re.split(r'[\r\n]+', nombres_raw) if n.strip()]
    cargos_list = [c.strip() for c in re.split(r'[\r\n]+', cargos_raw) if c.strip()]
    
    destinatarios = []
    max_len = max(len(nombres_list), len(cargos_list), 1)
    
    for i in range(max_len):
        nom = nombres_list[i] if i < len(nombres_list) else ""
        car = cargos_list[i] if i < len(cargos_list) else ""
        if nom or car:
            destinatarios.append({"nombre": nom, "cargo": car})
            
    if not destinatarios:
        destinatarios = [{"nombre": "", "cargo": ""}]
        
    tipo = "MULTIPLE" if len(destinatarios) > 1 else "SIMPLE"
    return destinatarios, tipo

def init_db():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    
    cur.execute('''
    CREATE TABLE IF NOT EXISTS proveidos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        numero INTEGER NOT NULL UNIQUE,
        codigo_formateado TEXT NOT NULL,
        anio INTEGER NOT NULL DEFAULT 2026,
        sufijo TEXT NOT NULL DEFAULT '-2026-UNHEVAL/EPG-D',
        tipo TEXT NOT NULL DEFAULT 'SIMPLE',
        destinatarios TEXT NOT NULL,
        destinatario_principal TEXT,
        cargo_principal TEXT,
        asunto TEXT NOT NULL,
        referencia TEXT,
        remitente_ref TEXT,
        registro TEXT,
        fecha_recibido TEXT,
        fecha_emision TEXT,
        lugar_emision TEXT DEFAULT 'Cayhuayna',
        disposicion TEXT,
        check_opciones TEXT,
        dni TEXT,
        titulo_art TEXT,
        estado TEXT DEFAULT 'EMITIDO',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )
    ''')
    
    cur.execute('''
    CREATE TABLE IF NOT EXISTS elevaciones (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        numero INTEGER NOT NULL UNIQUE,
        codigo_formateado TEXT NOT NULL,
        anio INTEGER NOT NULL DEFAULT 2026,
        sufijo TEXT NOT NULL DEFAULT '-2026-UNHEVAL/EPG-D',
        destinatario TEXT NOT NULL,
        asunto TEXT NOT NULL,
        referencia TEXT,
        remitente_ref TEXT,
        registro TEXT,
        fecha_recibido TEXT,
        fecha_emision TEXT,
        lugar_emision TEXT DEFAULT 'Cayhuayna',
        disposicion TEXT,
        check_opciones TEXT,
        doc_ref TEXT,
        estado TEXT DEFAULT 'EMITIDO',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )
    ''')
    
    cur.execute('''
    CREATE TABLE IF NOT EXISTS destinatarios_frecuentes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre TEXT UNIQUE NOT NULL,
        cargo TEXT,
        frecuencia INTEGER DEFAULT 1
    )
    ''')
    
    cur.execute('''
    CREATE TABLE IF NOT EXISTS configuracion (
        clave TEXT PRIMARY KEY,
        valor TEXT
    )
    ''')
    
    conn.commit()
    conn.close()

def migrate():
    print(f"Abriendo Excel: {EXCEL_PATH} ...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    
    # 1. Migrar REDACTAR PROV.
    ws_prov = wb['REDACTAR PROV.']
    count_prov = 0
    freq_dest = {}
    
    for r in range(2, 4000):
        c_num = ws_prov.cell(r, 1).value
        c_a = ws_prov.cell(r, 2).value
        c_cargo = ws_prov.cell(r, 3).value
        c_asunto = ws_prov.cell(r, 4).value
        c_ref = ws_prov.cell(r, 5).value
        c_ta = ws_prov.cell(r, 6).value
        c_reg = ws_prov.cell(r, 7).value
        c_fec_rec = ws_prov.cell(r, 8).value
        c_fec_emi = ws_prov.cell(r, 9).value
        c_disp = ws_prov.cell(r, 10).value
        c_dni = ws_prov.cell(r, 11).value
        c_titulo = ws_prov.cell(r, 12).value
        
        # Ignorar filas vacías
        if not clean_str(c_a) and not clean_str(c_asunto) and not clean_str(c_ref):
            continue
            
        num_str = clean_str(c_num)
        if not num_str:
            continue
            
        try:
            num_int = int(re.sub(r'\D', '', num_str))
        except:
            num_int = r - 1
            
        cod_formateado = f"{num_int:04d}" if num_int < 1000 else f"{num_int}"
        
        dest_list, tipo = parse_destinatarios(c_a, c_cargo)
        dest_principal = dest_list[0]["nombre"] if dest_list else ""
        cargo_principal = dest_list[0]["cargo"] if dest_list else ""
        
        for d in dest_list:
            nom = d["nombre"].strip()
            car = d["cargo"].strip()
            if nom:
                if nom not in freq_dest:
                    freq_dest[nom] = {"cargo": car, "count": 1}
                else:
                    freq_dest[nom]["count"] += 1
                    if car and not freq_dest[nom]["cargo"]:
                        freq_dest[nom]["cargo"] = car

        cur.execute('''
            INSERT OR REPLACE INTO proveidos (
                numero, codigo_formateado, anio, sufijo, tipo, destinatarios,
                destinatario_principal, cargo_principal, asunto, referencia,
                remitente_ref, registro, fecha_recibido, fecha_emision,
                disposicion, check_opciones, dni, titulo_art
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            num_int,
            cod_formateado,
            2026,
            '-2026-UNHEVAL/EPG-D',
            tipo,
            json.dumps(dest_list, ensure_ascii=False),
            dest_principal,
            cargo_principal,
            clean_str(c_asunto),
            clean_str(c_ref),
            clean_str(c_ta),
            clean_str(c_reg),
            format_date(c_fec_rec),
            format_date(c_fec_emi),
            clean_str(c_disp),
            json.dumps([]),
            clean_str(c_dni),
            clean_str(c_titulo)
        ))
        count_prov += 1
        
    print(f"-> Proveídos válidos migrados: {count_prov}")
    
    # 2. Migrar ELEVACIONES
    ws_elev = wb['ELEVACIONES']
    count_elev = 0
    
    for r in range(2, 1000):
        c_num = ws_elev.cell(r, 1).value
        c_a = ws_elev.cell(r, 2).value
        c_asunto = ws_elev.cell(r, 3).value
        c_ref = ws_elev.cell(r, 4).value
        c_remitente = ws_elev.cell(r, 5).value
        c_reg = ws_elev.cell(r, 6).value
        c_fec_rec = ws_elev.cell(r, 7).value
        c_fec_emi = ws_elev.cell(r, 8).value
        c_doc_ref = ws_elev.cell(r, 10).value
        
        # Ignorar filas sin asunto y sin referencia
        if not clean_str(c_asunto) and not clean_str(c_ref):
            continue
            
        num_str = clean_str(c_num)
        if not num_str:
            continue
            
        try:
            num_int = int(re.sub(r'\D', '', num_str))
        except:
            num_int = r - 1
            
        cod_formateado = f"{num_int:04d}"
        
        cur.execute('''
            INSERT OR REPLACE INTO elevaciones (
                numero, codigo_formateado, anio, sufijo, destinatario,
                asunto, referencia, remitente_ref, registro,
                fecha_recibido, fecha_emision, disposicion, check_opciones, doc_ref
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            num_int,
            cod_formateado,
            2026,
            '-2026-UNHEVAL/EPG-D',
            clean_str(c_a),
            clean_str(c_asunto),
            clean_str(c_ref),
            clean_str(c_remitente),
            clean_str(c_reg),
            format_date(c_fec_rec),
            format_date(c_fec_emi),
            "",
            json.dumps([]),
            clean_str(c_doc_ref)
        ))
        count_elev += 1
        
    print(f"-> Elevaciones válidas migradas: {count_elev}")
    
    # 3. Guardar destinatarios frecuentes
    for nom, data in freq_dest.items():
        cur.execute('''
            INSERT OR REPLACE INTO destinatarios_frecuentes (nombre, cargo, frecuencia)
            VALUES (?, ?, ?)
        ''', (nom, data["cargo"], data["count"]))
        
    # 4. Configuración inicial
    cur.execute("INSERT OR REPLACE INTO configuracion (clave, valor) VALUES ('anio_activo', '2026')")
    cur.execute("INSERT OR REPLACE INTO configuracion (clave, valor) VALUES ('institucion', 'UNIVERSIDAD NACIONAL HERMILIO VALDIZÁN')")
    cur.execute("INSERT OR REPLACE INTO configuracion (clave, valor) VALUES ('dependencia', 'ESCUELA DE POSGRADO')")
    cur.execute("INSERT OR REPLACE INTO configuracion (clave, valor) VALUES ('lugar_defecto', 'Cayhuayna')")
    cur.execute("INSERT OR REPLACE INTO configuracion (clave, valor) VALUES ('sufijo_defecto', '-2026-UNHEVAL/EPG-D')")
    
    conn.commit()
    conn.close()
    print("¡Base de datos limpia y lista con éxito!")

if __name__ == '__main__':
    if os.path.exists(DB_PATH):
        os.remove(DB_PATH)
    init_db()
    migrate()
